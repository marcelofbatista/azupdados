"""Gera planilhas SINTÉTICAS no layout dos indicadores educacionais do INEP por escola. Valores inventados.

Reproduzem de propósito a armadilha 12:
  - ATU 2017: rótulos de código e nome da escola/município trocados no cabeçalho técnico;
  - Rendimento 2018 e ATU 2017: nomes técnicos deslocados em relação aos rótulos (o dado segue o rótulo);
  - as duas gerações de nomes técnicos (tap_f14 × 1_cat_fun_ai, mird × edu_bas_cat_0...);
  - zip com .xlsx + .ods + md5 oficial; '--' como ausente; nota de rodapé no fim.

Escolas: 99000001 (só anos iniciais e infantil), 99000002 (fundamental completo e médio), 99000004 (só médio)
— as mesmas do Censo sintético — e 99000099 (não existe no Censo: testa o left join com dim_escola).

Uso: python data/fixtures/indicadores/gerar_fixtures.py
"""

import hashlib
import io
import zipfile
from pathlib import Path

import openpyxl

AQUI = Path(__file__).parent
ESCOLAS = {  # código: etapas oferecidas
    "99000001": {"inf", "ai"},
    "99000002": {"ai", "af", "em"},
    "99000004": {"em"},
    "99000099": {"ai"},
}
ID_ROTULOS = ["Ano", "Região", "UF", "Código do Município", "Nome do Município", "Código da Escola",
              "Nome da Escola", "Localização", "Dependência Administrativa"]
ID_NOVO = ["NU_ANO_CENSO", "NO_REGIAO", "SG_UF", "CO_MUNICIPIO", "NO_MUNICIPIO", "CO_ENTIDADE", "NO_ENTIDADE",
           "NO_CATEGORIA", "NO_DEPENDENCIA"]
ID_ANTIGO = ["NU_ANO_CENSO", "NO_REGIAO", "SG_UF", "CO_MUNICIPIO", "NO_MUNICIPIO", "CO_ENTIDADE", "NO_ENTIDADE",
             "TIPOLOCA", "DEPENDAD"]
ID_2015 = ["Ano", "NO_REGIAO", "SIGLA", "PK_COD_MUNICIPIO", "NO_MUNICIPIO", "PK_COD_ENTIDADE", "NO_ENTIDADE",
           "TIPOLOCA", "Dependad"]

# (rótulo, sufixo antigo, sufixo novo, oferta exigida)
EF = [("Total", "fun", "fun", {"ai", "af"}), ("Anos Iniciais", "f14", "fun_ai", {"ai"}),
      ("Anos Finais", "f58", "fun_af", {"af"})]
EF += [(f"{i}º Ano", f"f0{i - 1}", f"fun_0{i}", {"ai"} if i <= 5 else {"af"}) for i in range(1, 10)]
EM = [("Total", "med", "med", {"em"})] + [(f"{i}ª série", f"m0{i}", f"med_0{i}", {"em"}) for i in range(1, 5)]
EM += [("Não-Seriado", "mns", "med_ns", {"em"})]


def oferece(cod, exigido):
    return bool(ESCOLAS[cod] & exigido)


def valor(cod, i, base):
    return round(base + (int(cod[-1]) + i) % 5, 1)


def id_valores(cod, ano, trocar=False):
    v = [ano, "Centro-Oeste", "DF", "5300108", "Brasília", cod, f"ESCOLA {cod}", "Urbana", "Estadual"]
    if trocar:  # como no arquivo real: os DADOS vêm código antes do nome
        pass
    return v


def planilha(aba, rotulos_niveis, nomes, linhas):
    """rotulos_niveis: lista de linhas de rótulos (cada uma com o mesmo tamanho de nomes; None = mesclada)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = aba
    for r in (["Ministério da Educação"], ["Instituto Nacional de Estudos e Pesquisas Educacionais"], [],
              [f"Indicador sintético — {aba}"]):
        ws.append(r)
    for r in rotulos_niveis:
        ws.append(r)
    ws.append(nomes)
    for r in linhas:
        ws.append(r)
    ws.append(["Fonte: INEP (dados sintéticos para teste)."])
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def zipar(pasta, nome_zip, interno, xlsx):
    (AQUI / pasta).mkdir(exist_ok=True)
    md5 = hashlib.md5(xlsx).hexdigest().upper()
    with zipfile.ZipFile(AQUI / pasta / nome_zip, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr(f"{interno}/{interno}.xlsx", xlsx)
        z.writestr(f"{interno}/{interno}.ods", b"mesmo conteudo em ods")
        z.writestr(f"{interno}/md5_{interno}.txt", f"{md5} *{interno}.xlsx\r\n")


def rendimento(ano, antigo, deslocar):
    grupos = [("Aprovação", "tap", 1), ("Reprovação", "tre", 2), ("Abandono", "tab", 3)]
    r1, r2, r3, nomes = list(ID_ROTULOS), [None] * 9, [None] * 9, list(ID_ANTIGO if antigo else ID_NOVO)
    blocos = []
    for g, pre, k in grupos:
        for nivel, etapas in (("Ensino Fundamental de 8 e 9 anos", EF), ("Ensino Médio", EM)):
            sufixos = [e[1] for e in etapas]
            if deslocar and nivel.startswith("Ensino Fundamental"):
                sufixos.remove("f04")
                sufixos.insert(2, "f04")  # defeito real (2018–2020): f04 logo depois de f14
            for j, (rot, _, novo, exig) in enumerate(etapas):
                r1.append(f"Taxa de {g}" if (nivel.startswith("Ensino F") and j == 0) else None)
                r2.append(nivel if j == 0 else None)
                r3.append(rot)
                nomes.append(f"{pre}_{sufixos[j]}" if antigo else f"{k}_CAT_{novo.upper()}")
                blocos.append((k, exig, j))
    linhas = []
    for cod in ESCOLAS:
        v = id_valores(cod, ano)
        for k, exig, j in blocos:
            if not oferece(cod, exig):
                v.append("--")
                continue
            ap, rp = valor(cod, j, 88), valor(cod, j, 4)
            v.append(ap if k == 1 else rp if k == 2 else round(100 - ap - rp, 1))
        linhas.append(v)
    return planilha("ESCOLAS", [r1, r2, r3], nomes, linhas)


def etapa_unica(ano, antigo, pre, aba, com_infantil, deslocar=False, trocar_id=False, base=20):
    etapas = []
    if com_infantil:
        etapas += [("Educação Infantil", "Total", "inf", "ed_inf", {"inf"}),
                   ("Educação Infantil", "Creche", "cre", "cre", {"inf"}),
                   ("Educação Infantil", "Pré-Escola", "pre", "pre", {"inf"})]
    etapas += [("Ensino Fundamental", r, a, n, e) for r, a, n, e in EF]
    etapas += [("Ensino Médio", r, a, n, e) for r, a, n, e in EM]
    sufixos = [e[2] for e in etapas]
    if deslocar:
        sufixos.remove("f04")
        sufixos.insert(sufixos.index("f14") + 1, "f04")
    ids = list(ID_ANTIGO if antigo else ID_NOVO)
    if trocar_id:  # cabeçalho na ordem de 2015, dados na ordem nova (ATU 2016–2018)
        ids[3], ids[4], ids[5], ids[6] = "NO_MUNICIPIO", "CO_MUNICIPIO", "NO_ENTIDADE", "CO_ENTIDADE"
    r1, r2, nomes = list(ID_ROTULOS), [None] * 9, ids
    anterior = None
    for i, (nivel, rot, _, novo, _) in enumerate(etapas):
        r1.append(nivel if nivel != anterior else None)
        anterior = nivel
        r2.append(rot)
        nomes.append(f"{pre}_{sufixos[i]}" if antigo else f"{novo.upper()}_CAT_0")
    linhas = []
    for cod in ESCOLAS:
        v = id_valores(cod, ano)
        v += [valor(cod, i, base) if oferece(cod, e[4]) else "--" for i, e in enumerate(etapas)]
        linhas.append(v)
    return planilha(aba, [r1, r2], nomes, linhas)


def categorias(ano, antigo, pre, aba, n_cat, palavra):
    etapas = [("Ensino Fundamental", "Total", "fun", "fun", {"ai", "af"}),
              ("Ensino Fundamental", "Anos Iniciais", "f14", "fun_ai", {"ai"}),
              ("Ensino Fundamental", "Anos Finais", "f58", "fun_af", {"af"}),
              ("Ensino Médio", "Total", "med", "med", {"em"})]
    ids = list(ID_2015 if ano == 2015 else ID_NOVO)
    r1, r2, r3, nomes = list(ID_ROTULOS), [None] * 9, [None] * 9, ids
    for nivel, rot, ant, novo, _ in etapas:
        for c in range(1, n_cat + 1):
            r1.append(nivel if c == 1 else None)
            r2.append(rot if c == 1 else None)
            r3.append(f"{palavra} {c}")
            nomes.append(f"{pre}_{ant}{c}" if antigo else f"{novo.upper()}_CAT_{c}")
    linhas = []
    for cod in ESCOLAS:
        v = id_valores(cod, ano)
        for nivel, rot, ant, novo, exig in etapas:
            if not oferece(cod, exig):
                v += ["--"] * n_cat
                continue
            partes = [round(100 / n_cat + (c - 1) * 0.5 - (n_cat - 1) * 0.25, 1) for c in range(1, n_cat + 1)]
            partes[-1] = round(100 - sum(partes[:-1]), 1)
            v += partes
        linhas.append(v)
    return planilha(aba, [r1, r2, r3], nomes, linhas)


def escola_unico(ano, antigo, nome_ant, nome_novo, rotulo, valores):
    ids = list(ID_2015 if antigo else ID_NOVO)
    linhas = [id_valores(cod, ano) + [valores(cod)] for cod in ESCOLAS]
    return planilha("ESCOLAS", [ID_ROTULOS + [rotulo]], ids + [nome_ant if antigo else nome_novo], linhas)


def inse(ano):
    wb = openpyxl.Workbook()
    ws = wb.active
    if ano == 2015:
        ws.title = "INSE_2015"
        for _ in range(7):
            ws.append([])
        ws.append(["Indicador de Nível Socioeconômico das Escolas de Educação Básica - Inse 2015"])
        ws.append(["*Mais informações sobre a metodologia na nota técnica"])
        ws.append(["CO_ESCOLA", "NOME_ESCOLA", "CO_UF", "NOME_UF", "CO_MUNICIPIO", "NOME_MUNICIPIO", "ID_AREA",
                   "TP_DEPENDENCIA", "TP_LOCALIZACAO", "QTD_ALUNOS_INSE", "INSE_VALOR_ABSOLUTO",
                   "INSE_CLASSIFICACAO"])
        for i, cod in enumerate(ESCOLAS):
            ws.append([int(cod), f"ESCOLA {cod}", 53, "Distrito Federal", 5300108, "Brasília", 1, 2, 1,
                       30 + i, 50.5 + i, f"Grupo {3 + i % 3}"])
    else:
        ws.title = f"INSE_ESC_{ano}" if ano >= 2021 else f"INSE_{ano}"
        if ano >= 2021:
            cab = ["NU_ANO_SAEB", "CO_UF", "SG_UF", "NO_UF", "CO_MUNICIPIO", "NO_MUNICIPIO", "ID_ESCOLA",
                   "NO_ESCOLA", "TP_TIPO_REDE", "TP_LOCALIZACAO", "TP_CAPITAL", "QTD_ALUNOS_INSE", "MEDIA_INSE",
                   "INSE_CLASSIFICACAO"]
        else:
            cab = ["CO_ESCOLA", "NOME_ESCOLA", "CO_UF", "NOME_UF", "CO_MUNICIPIO", "NOME_MUNICIPIO", "ID_AREA",
                   "TP_DEPENDENCIA", "TP_LOCALIZACAO", "QTD_ALUNOS_INSE", "INSE_VALOR_ABSOLUTO",
                   "INSE_CLASSIFICACAO"]
        ws.append(cab + [f"PC_NIVEL_{n}" for n in range(1, 9)])
        for i, cod in enumerate(ESCOLAS):
            pcs = [0, 2.5, 10, 30, 40, 12.5, 5, 0]
            base = ([ano, 53, "DF", "Distrito Federal", 5300108, "Brasília", int(cod), f"ESCOLA {cod}", 2, 1, 1]
                    if ano >= 2021 else
                    [int(cod), f"ESCOLA {cod}", 53, "Distrito Federal", 5300108, "Brasília", 1, 2, 1])
            ws.append(base + [40 + i, 4.8 + i / 10, ["Nível III", "Nível IV", "Nível V", "Nível VI"][i]] + pcs)
        d = wb.create_sheet("Dicionário")
        d.append(["Nome", "Descrição"])
        d.append(["CO_ESCOLA", "Código da Escola no Censo Escolar"])
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def main():
    zipar("rendimento", "TX_REND_ESCOLAS_2018.zip", "TX_REND_ESCOLAS_2018", rendimento(2018, True, True))
    zipar("rendimento", "tx_rend_escolas_2021.zip", "tx_rend_escolas_2021", rendimento(2021, False, False))
    zipar("tdi", "TDI_2016_ESCOLAS.zip", "TDI_ESCOLAS_2016", etapa_unica(2016, True, "tdi", "ESCOLAS", False))
    zipar("tdi", "TDI_2021_ESCOLAS.zip", "TDI_ESCOLAS_2021", etapa_unica(2021, False, "tdi", "ESCOLAS", False))
    zipar("atu", "ATU_2017_ESCOLAS.zip", "ATU_ESCOLAS_2017",
          etapa_unica(2017, True, "atu", "ESCOLA", True, deslocar=True, trocar_id=True))
    zipar("had", "HAD_2021_ESCOLAS.zip", "HAD_ESCOLAS_2021",
          etapa_unica(2021, False, "had", "ESCOLA", True, base=4))
    zipar("afd", "adequacao_formacao_docente_escolas_2015.zip", "AFD_ESCOLAS_2015",
          categorias(2015, True, "afd", "Ind. adeq. form. doc.", 5, "Grupo"))
    zipar("ied", "IED_2019_ESCOLAS.zip", "IED_ESCOLAS_2019", categorias(2019, False, "ied", "ESCOLAS", 6, "Nível"))
    zipar("ird", "regularidade_corpo_docente_escolas_2015.zip", "IRD_ESCOLAS_2015",
          escola_unico(2015, True, "MIRD", "EDU_BAS_CAT_0", "Média do Indicador de Regularidade do Docente (IRD)",
                       lambda c: 3.25 + int(c[-1]) / 10))
    zipar("icg", "ICG_2019_ESCOLAS.zip", "ICG_ESCOLAS_2019",
          escola_unico(2019, False, "ICG", "COMPLEX", "Nível de complexidade de gestão da escola3",
                       lambda c: f"Nível {1 + int(c[-1]) % 6}"))
    (AQUI / "inse").mkdir(exist_ok=True)
    for ano in (2019, 2023):
        (AQUI / "inse" / (f"INSE_{ano}_ESCOLAS.xlsx" if ano == 2019 else f"INSE_{ano}_escolas.xlsx")).write_bytes(inse(ano))
    with zipfile.ZipFile(AQUI / "inse" / "Indicador_Inse_2015.zip", "w") as z:
        z.writestr("INSE_2015.xlsx", inse(2015))
        z.writestr("nota_tecnica_indicador_escola_nivel_socioeconomico_2015.pdf", b"%PDF-1.4 sintetico")
    print("fixtures geradas em", AQUI)


if __name__ == "__main__":
    main()
